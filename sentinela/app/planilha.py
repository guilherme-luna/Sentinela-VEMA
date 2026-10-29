"""Saída para a planilha "01 - Controle de Acervo" usada pela vara.

Mantém o formato atual (uma aba por servidor + aba Unificada) e acrescenta as
colunas URGÊNCIA e VÍNCULOS. As colunas de controle de cada servidor ficam em
branco para preenchimento livre, como hoje. Na fase 2, a mesma escrita é feita
diretamente na planilha do Teams via Microsoft Graph.
"""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import SERVIDORES

COLUNAS = ["Quant.", "PROCESSOS", "SISTEMA", "CLASSE", "DATA DO RECEBIMENTO", "URGÊNCIA", "VÍNCULOS",
           "PROVIDÊNCIA", "ÚLTIMA ANÁLISE", "Tempo decorrido", "OBSERVAÇÃO"]
LARGURAS = [8, 28, 10, 34, 20, 34, 40, 40, 16, 15, 30]
ABA_UNIFICADA = "Unificada - Controle"

CABECALHO = PatternFill("solid", fgColor="D9D9D9")
CORES_URGENCIA = {1: "F4B6B6", 2: "F9D5A7", 3: "FBEFA6"}
BORDA = Border(*(Side(style="thin", color="999999"),) * 4)


def _linha(qtd: int, p: dict, com_servidor: bool) -> list:
    recebido = datetime.fromisoformat(p["data_entrada"]).date()
    urg = "; ".join(a["rotulo"] for a in p.get("alertas", []))
    vinc = "; ".join(
        f"{v['numero']} ({v['criterios'][0]['detalhe']})" for v in p.get("vinculos", [])[:3])
    linha = [qtd, p["numero"], "PJE", p["classe"].upper(), recebido, urg, vinc, "", "",
             (date.today() - recebido).days, ""]
    if com_servidor:
        linha.insert(2, p.get("servidor"))
    return linha


def _aba(wb: Workbook, titulo: str, processos: list[dict], com_servidor: bool = False):
    colunas, larguras = list(COLUNAS), list(LARGURAS)
    if com_servidor:
        colunas.insert(2, "SERVIDOR")
        larguras.insert(2, 14)
    desloc = 1 if com_servidor else 0
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = Font(bold=True), CABECALHO, BORDA
        c.alignment = Alignment(horizontal="center")
    for qtd, p in enumerate(processos, start=1):
        ws.append(_linha(qtd, p, com_servidor))
        r = ws.max_row
        for col in range(1, len(colunas) + 1):
            ws.cell(row=r, column=col).border = BORDA
        ws.cell(row=r, column=5 + desloc).number_format = "DD/MM/YYYY"
        if p.get("alertas"):
            nivel = min(a["nivel"] for a in p["alertas"])
            ws.cell(row=r, column=6 + desloc).fill = PatternFill("solid", fgColor=CORES_URGENCIA[nivel])
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def gerar(acervo: list[dict], destino: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for s in SERVIDORES:
        _aba(wb, s.nome, [p for p in acervo if p.get("servidor") == s.nome])
    _aba(wb, ABA_UNIFICADA, sorted(acervo, key=lambda p: p["data_entrada"]), com_servidor=True)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino
