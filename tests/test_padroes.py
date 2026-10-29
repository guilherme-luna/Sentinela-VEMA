"""Testes dos padrões decisórios e da decisão nos casos semelhantes."""
from openpyxl import load_workbook

from app import cnj, padroes, planilha, vinculos

MP = {"nome": "Ministério Público do Estado do Ceará", "documento": None, "polo": "ativo"}
EMPRESA = {"nome": "Lavanderia Teste Ltda", "documento": "07.384.912/0001-55", "polo": "passivo"}


def proc(seq, assunto="Poluição", decisao=None, entrada="2026-05-01T10:00", partes=None, sintese=None):
    return {
        "numero": cnj.montar(seq, 2026), "classe": "Ação Civil Pública", "assunto": assunto,
        "data_entrada": entrada, "partes": partes or [MP], "locais": [], "servidor": "Ana",
        "sintese": sintese or f"lançamento de efluentes na lagoa, caso {seq}",
        "decisao": {"tipo": decisao, "resumo": "resumo", "data": "2026-05-21T10:00"} if decisao else None,
    }


def test_resumir_aponta_a_decisao_predominante():
    relacionados = [proc(1, decisao="Liminar deferida"), proc(2, decisao="Liminar deferida"),
                    proc(3, decisao="TAC homologado"), proc(4)]
    r = padroes.resumir(relacionados)
    assert r["decididos"] == 3  # o caso sem decisão não entra na conta
    assert (r["predominante"], r["vezes"]) == ("Liminar deferida", 2)
    assert r["distribuicao"] == [{"tipo": "Liminar deferida", "n": 2}, {"tipo": "TAC homologado", "n": 1}]


def test_resumir_sem_decisoes_devolve_nada():
    assert padroes.resumir([proc(1), proc(2)]) is None
    assert padroes.resumir([]) is None


def test_quadro_por_assunto_conta_prazo_e_recentes():
    acervo = [proc(1, decisao="Liminar deferida", entrada="2026-05-01T10:00"),
              proc(2, decisao="TAC homologado", entrada="2026-05-11T10:00"),
              proc(3), proc(4, assunto="Crimes contra a Fauna", decisao="Transação penal")]
    quadro = {g["assunto"]: g for g in padroes.por_assunto(acervo)}
    poluicao = quadro["Poluição"]
    assert (poluicao["total"], poluicao["decididos"]) == (3, 2)
    assert poluicao["dias_ate_decisao"] == 15  # mediana entre 20 e 10 dias
    assert {r["numero"] for r in poluicao["recentes"]} == {cnj.montar(1, 2026), cnj.montar(2, 2026)}
    assert list(quadro)[0] == "Poluição"  # ordenado pelo tamanho do assunto


def test_caso_vinculado_traz_a_decisao_da_vara():
    anterior = proc(1, decisao="Liminar deferida", partes=[MP, EMPRESA])
    novo = proc(2, partes=[MP, EMPRESA])
    novo["decisao"] = None
    achados = vinculos.encontrar(novo, [anterior])
    assert achados["vinculos"][0]["decisao"]["tipo"] == "Liminar deferida"


def test_planilha_inicial_traz_providencia_registrada(tmp_path):
    destino = tmp_path / "planilha.xlsx"
    planilha.gerar([proc(1, decisao="Liminar deferida"), proc(2)], destino)
    ws = load_workbook(destino)["Ana"]
    assert ws["H2"].value == "Liminar deferida: resumo"
    assert ws["I2"].value is not None  # ÚLTIMA ANÁLISE = data da decisão
    assert ws["H3"].value in (None, "")  # sem decisão, a coluna fica livre para o servidor
