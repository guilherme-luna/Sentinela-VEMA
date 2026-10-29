"""Testes do fluxo completo: extração incremental, duplicados, aprovação, planilha e concorrência."""
import threading

import pytest
from openpyxl import load_workbook

from app import cnj
from app.repositorio import Repositorio
from app.servico import Sentinela

MP = {"nome": "Ministério Público do Estado do Ceará", "documento": None, "polo": "ativo"}


def proc(seq: int, quando: str, **extra) -> dict:
    return {
        "numero": cnj.montar(seq, 2026), "classe": "Ação Civil Pública", "competencia": "Cível",
        "assunto": "Poluição", "data_entrada": quando, "partes": [MP], "apenso_de": None,
        "metadados": {}, "locais": [], "sintese": f"fatos do processo {seq}", **extra,
    }


class FilaFake:
    """PJe simulado em memória: a fila cresce entre uma busca e outra."""

    def __init__(self, itens=()):
        self.itens = list(itens)

    def buscar_novos(self, desde):
        return [p for p in self.itens if desde is None or p["data_entrada"] > desde.isoformat()]

    def paginas_documento(self, p):
        return p.get("_paginas", [])


@pytest.fixture
def sistema(tmp_path):
    repo = Repositorio(tmp_path / "teste.db")
    fila = FilaFake()
    s = Sentinela(repo, fila, planilha_path=tmp_path / "planilha.xlsx")
    repo.set_estado("marco", "2026-10-01T00:00")
    return s, repo, fila


def test_extracao_incremental_nao_repete_o_que_ja_foi_aprovado(sistema):
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00"), proc(2, "2026-10-02T11:00")]
    assert s.executar_triagem()["novos"] == 2
    s.aprovar()
    fila.itens.append(proc(3, "2026-10-03T09:00"))
    r = s.executar_triagem()
    assert r["novos"] == 1 and r["duplicados"] == []


def test_duplicado_e_bloqueado_com_nome_do_servidor(sistema):
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    s.aprovar()
    dono = repo.servidor_de(cnj.montar(1, 2026))
    fila.itens.append(proc(1, "2026-10-04T10:00"))  # voltou para a fila (redistribuição)
    r = s.executar_triagem()
    assert r["novos"] == 0
    assert r["duplicados"][0]["motivo"] == f"Já distribuído para {dono}"


def test_marco_avanca_quando_so_ha_duplicados(sistema):
    """Antes da correção, o duplicado reaparecia em toda busca porque não havia o que aprovar."""
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    s.aprovar()
    fila.itens.append(proc(1, "2026-10-04T10:00"))
    assert len(s.executar_triagem()["duplicados"]) == 1
    assert s.executar_triagem()["duplicados"] == []


def test_marco_inclui_duplicado_lido_depois_do_ultimo_novo(sistema):
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    s.aprovar()
    fila.itens += [proc(2, "2026-10-05T08:00"), proc(1, "2026-10-05T09:00")]  # novo + duplicado depois
    r = s.executar_triagem()
    assert r["novos"] == 1 and len(r["duplicados"]) == 1
    s.aprovar()
    assert repo.get_estado("marco") == "2026-10-05T09:00"
    assert s.executar_triagem()["duplicados"] == []


def test_marco_so_avanca_apos_aprovacao(sistema):
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    assert repo.get_estado("marco") == "2026-10-01T00:00"
    assert s.executar_triagem()["novos"] == 1  # refazer a busca traz a mesma proposta


def test_numero_invalido_nao_entra(sistema):
    s, repo, fila = sistema
    ruim = proc(1, "2026-10-02T10:00")
    ruim["numero"] = ruim["numero"][:8] + ("1" if ruim["numero"][8] != "1" else "2") + ruim["numero"][9:]
    fila.itens = [ruim]
    r = s.executar_triagem()
    assert r["novos"] == 0 and r["invalidos"][0]["motivo"] == "Dígito verificador inválido"


def test_aprovar_preserva_anotacoes_do_servidor_na_planilha(sistema):
    """Antes da correção, a planilha era recriada e apagava as colunas preenchidas pelos servidores."""
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    s.aprovar()
    caminho = s.planilha_path
    wb = load_workbook(caminho)
    servidor = repo.servidor_de(cnj.montar(1, 2026))
    wb[servidor]["H2"] = "Minutado. À delegacia."
    wb.save(caminho)

    fila.itens.append(proc(2, "2026-10-03T10:00"))
    s.executar_triagem()
    s.aprovar()

    wb = load_workbook(caminho)
    assert wb[servidor]["H2"].value == "Minutado. À delegacia."
    unificada = [r for r in wb["Unificada - Controle"].iter_rows(min_row=2, values_only=True)]
    assert [r[0] for r in unificada] == [1, 2]  # Quant. continua a sequência
    assert {r[1] for r in unificada} == {cnj.montar(1, 2026), cnj.montar(2, 2026)}


def test_planilha_aberta_nao_grava_nada(sistema, monkeypatch):
    from app import planilha

    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()

    def bloqueada(*_a, **_k):
        raise planilha.PlanilhaEmUso("aberta no Excel")

    monkeypatch.setattr(planilha, "acrescentar", bloqueada)
    with pytest.raises(planilha.PlanilhaEmUso):
        s.aprovar()
    assert len(repo.listar("proposta")) == 1  # nada aprovado; pode tentar de novo
    assert repo.get_estado("marco") == "2026-10-01T00:00"


def test_reatribuir_so_em_proposta(sistema):
    s, repo, fila = sistema
    fila.itens = [proc(1, "2026-10-02T10:00")]
    s.executar_triagem()
    n = cnj.montar(1, 2026)
    s.reatribuir(n, "Elisa")
    assert repo.servidor_de(n) == "Elisa"
    s.aprovar()
    with pytest.raises(ValueError):
        s.reatribuir(n, "Ana")


def test_repositorio_suporta_acesso_concorrente(tmp_path):
    """Antes da correção, 6 threads usando a mesma conexão geravam InterfaceError."""
    repo = Repositorio(tmp_path / "concorrencia.db")
    repo.salvar_varios([proc(i, "2026-10-02T10:00") for i in range(1, 30)], "acervo")
    erros = []

    def usar():
        try:
            for _ in range(150):
                repo.listar("acervo")
                repo.auditar("teste", "x")
        except Exception as e:  # noqa: BLE001
            erros.append(e)

    threads = [threading.Thread(target=usar) for _ in range(6)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert erros == []
