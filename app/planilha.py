"""Saída para a planilha "01 - Controle de Acervo" usada pela vara.

Mantém o formato atual (uma aba por servidor + aba Unificada) e acrescenta as
colunas URGÊNCIA e VÍNCULOS. As colunas de controle de cada servidor
(PROVIDÊNCIA, ÚLTIMA ANÁLISE, OBSERVAÇÃO) são do servidor: a aprovação só
ACRESCENTA linhas novas e nunca reescreve o que já está na planilha. Na fase 2,
a mesma operação é feita na planilha do Teams via Microsoft Graph.
"""
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .config import SERVIDORES

COLUNAS = ["Quant.", "PROCESSOS", "SISTEMA", "CLASSE", "DATA DO RECEBIMENTO", "URGÊNCIA", "VÍNCULOS",
           "PROVIDÊNCIA", "ÚLTIMA ANÁLISE", "Tempo decorrido", "OBSERVAÇÃO"]
LARGURAS = [8, 28, 10, 34, 20, 34, 40, 40, 16, 15, 30]
ABA_UNIFICADA = "Unificada - Controle"

CABECALHO = PatternFill("solid", fgColor="D9D9D9")
CORES_URGENCIA = {1: "F4B6B6", 2: "F9D5A7", 3: "FBEFA6"}
BORDA = Border(*(Side(style="thin", color="999999"),) * 4)


class PlanilhaEmUso(RuntimeError):
    """O arquivo está aberto em outro programa (ex.: Excel) e não pode ser gravado."""


def _colunas(com_servidor: bool) -> tuple[list[str], list[int]]:
    colunas, larguras = list(COLUNAS), list(LARGURAS)
    if com_servidor:
        colunas.insert(2, "SERVIDOR")
        larguras.insert(2, 14)
    return colunas, larguras


def _nova_aba(wb: Workbook, titulo: str, com_servidor: bool) -> Worksheet:
    colunas, larguras = _colunas(com_servidor)
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        c = ws.cell(row=1, column=i)
        c.font, c.fill, c.border = Font(bold=True), CABECALHO, BORDA
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    return ws


def _proxima_quantidade(ws: Worksheet) -> int:
    ultimo = 0
    for (valor,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if isinstance(valor, (int, float)):
            ultimo = max(ultimo, int(valor))
    return ultimo + 1


def _escrever(ws: Worksheet, p: dict, com_servidor: bool):
    recebido = datetime.fromisoformat(p["data_entrada"]).date()
    urg = "; ".join(a["rotulo"] for a in p.get("alertas", []))
    vinc = "; ".join(f"{v['numero']} ({v['criterios'][0]['detalhe']})" for v in p.get("vinculos", [])[:3])
    decisao = p.get("decisao") or {}
    analise = datetime.fromisoformat(decisao["data"]).date() if decisao else ""
    providencia = f"{decisao['tipo']}: {decisao['resumo']}" if decisao else ""
    linha = [_proxima_quantidade(ws), p["numero"], "PJE", p["classe"].upper(), recebido, urg, vinc, providencia,
             analise, (date.today() - recebido).days, ""]
    if com_servidor:
        linha.insert(2, p.get("servidor"))
    ws.append(linha)

    desloc = 1 if com_servidor else 0
    r = ws.max_row
    for col in range(1, len(linha) + 1):
        ws.cell(row=r, column=col).border = BORDA
    ws.cell(row=r, column=5 + desloc).number_format = "DD/MM/YYYY"
    ws.cell(row=r, column=9 + desloc).number_format = "DD/MM/YYYY"
    if p.get("alertas"):
        nivel = min(a["nivel"] for a in p["alertas"])
        ws.cell(row=r, column=6 + desloc).fill = PatternFill("solid", fgColor=CORES_URGENCIA[nivel])
    ws.auto_filter.ref = ws.dimensions


def _salvar(wb: Workbook, destino: Path):
    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(destino)
    except PermissionError as e:
        raise PlanilhaEmUso(f"A planilha '{destino.name}' está aberta em outro programa. "
                            "Feche-a e aprove novamente.") from e


def gerar(acervo: list[dict], destino: Path) -> Path:
    """Cria a planilha do zero com todo o acervo (carga inicial)."""
    wb = Workbook()
    wb.remove(wb.active)
    for s in SERVIDORES:
        ws = _nova_aba(wb, s.nome, com_servidor=False)
        for p in acervo:
            if p.get("servidor") == s.nome:
                _escrever(ws, p, com_servidor=False)
    ws = _nova_aba(wb, ABA_UNIFICADA, com_servidor=True)
    for p in sorted(acervo, key=lambda p: p["data_entrada"]):
        _escrever(ws, p, com_servidor=True)
    _salvar(wb, destino)
    return destino


def acrescentar(novos: list[dict], destino: Path, acervo_completo: list[dict]) -> Path:
    """Acrescenta os processos aprovados sem tocar nas linhas existentes.

    Se a planilha ainda não existe, cria com todo o acervo.
    """
    if not destino.exists():
        return gerar(acervo_completo, destino)
    wb = load_workbook(destino)
    for p in sorted(novos, key=lambda p: p["data_entrada"]):
        servidor = p.get("servidor") or "Sem servidor"
        ws = wb[servidor] if servidor in wb.sheetnames else _nova_aba(wb, servidor, com_servidor=False)
        _escrever(ws, p, com_servidor=False)
        uni = wb[ABA_UNIFICADA] if ABA_UNIFICADA in wb.sheetnames else _nova_aba(wb, ABA_UNIFICADA, com_servidor=True)
        _escrever(uni, p, com_servidor=True)
    _salvar(wb, destino)
    return destino
